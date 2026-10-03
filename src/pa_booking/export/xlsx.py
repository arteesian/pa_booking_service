"""Сборка xlsx из строк — общая для записей и библиотеки.

Даты и время пишутся значениями Excel (сортировка и фильтры в Excel работают),
формат отображения — русский: ``ДД.ММ.ГГГГ`` и ``ЧЧ:ММ``.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATE_FORMAT = "DD.MM.YYYY"
TIME_FORMAT = "HH:MM"
_MAX_WIDTH = 60


def _display_len(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, date | time):
        return 10
    return len(str(value))


def build_xlsx(title: str, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> bytes:
    """Один лист: жирная шапка, закреплённая первая строка, ширина колонок по содержимому."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # у нового Workbook активный лист есть всегда
    ws.title = title[:31]  # ограничение Excel на имя листа
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    widths = [len(h) for h in headers]
    for row in rows:
        ws.append(list(row))
        for i, value in enumerate(row):
            widths[i] = max(widths[i], _display_len(value))
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, datetime | date):
                cell.number_format = DATE_FORMAT
            elif isinstance(cell.value, time):
                cell.number_format = TIME_FORMAT
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, _MAX_WIDTH)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
