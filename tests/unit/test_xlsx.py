from __future__ import annotations

from datetime import date, time
from io import BytesIO

from openpyxl import load_workbook

from pa_booking.export.xlsx import build_xlsx


def _sheet(data: bytes):  # type: ignore[no-untyped-def]
    return load_workbook(BytesIO(data)).active


def test_header_bold_and_rows_in_order() -> None:
    ws = _sheet(
        build_xlsx(
            "Записи",
            ["Дата", "Время", "ФИО"],
            [
                (date(2026, 10, 5), time(16, 0), "Иванов Иван"),
                (date(2026, 10, 6), time(9, 30), None),
            ],
        )
    )
    assert ws.title == "Записи"
    assert [c.value for c in ws[1]] == ["Дата", "Время", "ФИО"]
    assert all(c.font.bold for c in ws[1])
    assert ws.max_row == 3
    assert ws["C3"].value is None


def test_dates_and_times_are_excel_values_with_russian_format() -> None:
    ws = _sheet(build_xlsx("Записи", ["Дата", "Время"], [(date(2026, 10, 5), time(16, 0))]))
    assert ws["A2"].is_date
    assert ws["A2"].value.date() == date(2026, 10, 5)
    assert ws["A2"].number_format == "DD.MM.YYYY"
    assert ws["B2"].value == time(16, 0)
    assert ws["B2"].number_format == "HH:MM"


def test_column_width_fits_longest_value() -> None:
    long_name = "Константинопольский Константин Константинович"
    ws = _sheet(build_xlsx("Записи", ["ФИО"], [(long_name,)]))
    assert ws.column_dimensions["A"].width >= len(long_name)


def test_empty_rows_give_header_only() -> None:
    ws = _sheet(build_xlsx("Записи", ["Дата"], []))
    assert ws.max_row == 1
