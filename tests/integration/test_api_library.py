from __future__ import annotations

import threading
import uuid
from datetime import date, datetime
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from pa_booking.db.models import DirectoryEmployee, LibraryBook
from pa_booking.domain.library import text_extended, text_loaned, text_returned
from pa_booking.domain.moscow import MSK
from pa_booking.notify.botx import FakeNotifier
from tests.integration.conftest import ApiEnv, headers

pytestmark = pytest.mark.db

# «сейчас» по умолчанию — 05.10.2026 10:00 МСК (фикстура api)
USER = uuid.UUID("11111111-1111-1111-1111-111111111111")
OTHER = uuid.UUID("22222222-2222-2222-2222-222222222222")
LIB = uuid.UUID("44444444-4444-4444-4444-444444444444")
TITLE = "Мастер и Маргарита"


def day(d: int, h: int = 10) -> datetime:
    return datetime(2026, 10, d, h, tzinfo=MSK)


def add_book(db: Session, title: str = TITLE, genre: str = "Роман") -> int:
    book = LibraryBook(genre=genre, author="Булгаков", title=title, description="—")
    db.add(book)
    db.commit()
    return book.id


def add_name(db: Session, employee_id: uuid.UUID, name: str) -> None:
    db.add(DirectoryEmployee(employee_id=employee_id, full_name=name))
    db.commit()


def lib_headers() -> dict[str, str]:
    return headers(user_id=LIB, roles="librarian")


def loan(api: ApiEnv, book_id: int, user: uuid.UUID = USER) -> httpx.Response:
    return api.client.post(f"/library/books/{book_id}/loan", headers=headers(user_id=user))


def act(api: ApiEnv, loan_id: int, action: str, user: uuid.UUID = USER) -> httpx.Response:
    return api.client.post(f"/library/loans/{loan_id}/{action}", headers=headers(user_id=user))


def catalog(api: ApiEnv) -> list[dict[str, object]]:
    r = api.client.get("/library/books", headers=headers(user_id=USER))
    assert r.status_code == 200
    result: list[dict[str, object]] = r.json()
    return result


# --- роли ---


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("POST", "/library/admin/books"),
        ("PATCH", "/library/admin/books/1"),
        ("DELETE", "/library/admin/books/1"),
        ("GET", "/library/admin/loans?state=active"),
        ("POST", "/library/admin/loans/1/return"),
        ("GET", "/library/admin/export"),
    ],
)
def test_admin_endpoints_require_librarian(api: ApiEnv, method: str, path: str) -> None:
    r = api.client.request(
        method,
        path,
        json={"genre": "Роман", "author": "А", "title": "Б", "description": "В"},
        headers=headers(user_id=USER, roles="operator,admin,psychologist"),
    )
    assert r.status_code == 403
    assert r.json()["code"] == "forbidden"


# --- каталог и бронь ---


def test_loan_free_book_for_a_week_and_notify(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    book_id = add_book(db_session)

    r = loan(api, book_id)

    assert r.status_code == 201
    body = r.json()
    assert (body["starts_on"], body["due_on"], body["overdue"]) == (
        "2026-10-05",
        "2026-10-12",
        False,
    )
    assert body["book"]["id"] == book_id
    assert api.notifier.sent == [text_loaned("Иванов Иван", TITLE, date(2026, 10, 12))]
    assert [(b["id"], b["available"]) for b in catalog(api)] == [(book_id, False)]


def test_loan_taken_book_is_409_without_notification(api: ApiEnv, db_session: Session) -> None:
    book_id = add_book(db_session)
    assert loan(api, book_id, user=OTHER).status_code == 201

    r = loan(api, book_id)

    assert r.status_code == 409
    assert r.json()["code"] == "book_unavailable"
    assert len(api.notifier.sent) == 1


def test_parallel_loans_of_one_book(api: ApiEnv, db_session: Session) -> None:
    book_id = add_book(db_session)
    barrier = threading.Barrier(2)
    results: list[httpx.Response] = []

    def worker(user: uuid.UUID) -> None:
        barrier.wait()
        results.append(loan(api, book_id, user=user))

    threads = [threading.Thread(target=worker, args=(u,)) for u in (USER, OTHER)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert sorted(r.status_code for r in results) == [201, 409]


def test_catalog_and_genres(api: ApiEnv, db_session: Session) -> None:
    taken = add_book(db_session, "Альфа")
    free = add_book(db_session, "Бета")
    add_book(db_session, "Гамма", genre="Детектив")
    loan(api, taken)

    r = api.client.get("/library/books", params={"genre": "Роман"}, headers=headers(user_id=USER))
    genres = api.client.get("/library/genres", headers=headers(user_id=USER))

    assert [(b["id"], b["available"]) for b in r.json()] == [(free, True), (taken, False)]
    assert genres.json() == ["Детектив", "Роман"]


def test_notify_failure_keeps_loan(api: ApiEnv, db_session: Session) -> None:
    api.notifier = FakeNotifier(fail=True)
    assert loan(api, add_book(db_session)).status_code == 201
    assert len(api.client.get("/library/loans/my", headers=headers(user_id=USER)).json()) == 1


# --- продление и возврат ---


def test_extend_adds_week_to_due_date(api: ApiEnv, db_session: Session) -> None:
    loan_id = loan(api, add_book(db_session)).json()["id"]
    api.clock.now = day(12, 23)  # в день срока ещё можно

    r = act(api, loan_id, "extend")

    assert r.status_code == 200
    assert r.json()["due_on"] == "2026-10-19"
    assert api.notifier.sent[-1] == text_extended(str(USER), TITLE, date(2026, 10, 19))


def test_extend_overdue_is_409(api: ApiEnv, db_session: Session) -> None:
    loan_id = loan(api, add_book(db_session)).json()["id"]
    api.clock.now = day(13)

    r = act(api, loan_id, "extend")

    assert r.status_code == 409
    assert r.json()["code"] == "loan_overdue"


def test_foreign_loan_is_404(api: ApiEnv, db_session: Session) -> None:
    loan_id = loan(api, add_book(db_session), user=OTHER).json()["id"]
    for action in ("extend", "return"):
        r = act(api, loan_id, action)
        assert r.status_code == 404
        assert r.json()["code"] == "not_found"


def test_return_overdue_loan_frees_book(api: ApiEnv, db_session: Session) -> None:
    book_id = add_book(db_session)
    loan_id = loan(api, book_id).json()["id"]
    api.clock.now = day(20)
    my = api.client.get("/library/loans/my", headers=headers(user_id=USER)).json()
    assert [(m["id"], m["overdue"]) for m in my] == [(loan_id, True)]  # Д-3: видна

    r = act(api, loan_id, "return")

    assert r.status_code == 200
    assert r.json()["returned_at"] is not None
    assert api.notifier.sent[-1] == text_returned(str(USER), TITLE, by_librarian=False)
    assert [b["available"] for b in catalog(api)] == [True]
    assert act(api, loan_id, "return").status_code == 404


def test_librarian_marks_return(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    loan_id = loan(api, add_book(db_session)).json()["id"]

    r = api.client.post(f"/library/admin/loans/{loan_id}/return", headers=lib_headers())

    assert r.status_code == 200
    assert r.json()["returned_by_librarian"] is True
    assert r.json()["full_name"] == "Иванов Иван"
    assert api.notifier.sent[-1] == text_returned("Иванов Иван", TITLE, by_librarian=True)
    again = api.client.post(f"/library/admin/loans/{loan_id}/return", headers=lib_headers())
    assert again.status_code == 404


def test_admin_loans_split_active_and_overdue(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    old = loan(api, add_book(db_session, "Альфа")).json()["id"]  # срок 12.10
    api.clock.now = day(10)
    fresh = loan(api, add_book(db_session, "Бета"), user=OTHER).json()["id"]  # срок 17.10
    api.clock.now = day(13)

    def ids(state: str) -> list[tuple[object, ...]]:
        r = api.client.get("/library/admin/loans", params={"state": state}, headers=lib_headers())
        assert r.status_code == 200
        return [(x["id"], x["full_name"], x["overdue"]) for x in r.json()]

    assert ids("active") == [(fresh, str(OTHER), False)]
    assert ids("overdue") == [(old, "Иванов Иван", True)]


# --- каталог библиотекаря ---


def test_admin_add_and_update_book(api: ApiEnv) -> None:
    r = api.client.post(
        "/library/admin/books",
        json={"genre": " Роман ", "author": "Булгаков", "title": TITLE, "description": "—"},
        headers=lib_headers(),
    )
    assert r.status_code == 201
    book_id = r.json()["id"]
    assert r.json()["genre"] == "Роман"

    r = api.client.patch(
        f"/library/admin/books/{book_id}",
        json={"title": "Белая гвардия", "author": None},
        headers=lib_headers(),
    )
    assert r.status_code == 200
    assert (r.json()["title"], r.json()["author"]) == ("Белая гвардия", "Булгаков")


def test_admin_add_book_rejects_empty_field(api: ApiEnv) -> None:
    r = api.client.post(
        "/library/admin/books",
        json={"genre": "Роман", "author": "  ", "title": TITLE, "description": "—"},
        headers=lib_headers(),
    )
    assert r.status_code == 422


def test_admin_remove_book_on_loan_is_409(api: ApiEnv, db_session: Session) -> None:
    book_id = add_book(db_session)
    loan(api, book_id)

    r = api.client.delete(f"/library/admin/books/{book_id}", headers=lib_headers())

    assert r.status_code == 409
    assert r.json()["code"] == "book_on_loan"


def test_admin_remove_free_book(api: ApiEnv, db_session: Session) -> None:
    book_id = add_book(db_session)

    assert (
        api.client.delete(f"/library/admin/books/{book_id}", headers=lib_headers()).status_code
        == 204
    )

    assert catalog(api) == []
    assert api.client.get("/library/genres", headers=headers(user_id=USER)).json() == []
    assert loan(api, book_id).status_code == 404
    r = api.client.patch(
        f"/library/admin/books/{book_id}", json={"title": "X"}, headers=lib_headers()
    )
    assert r.status_code == 404


# --- выгрузка ---


def test_export_history(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    first = loan(api, add_book(db_session, "Альфа")).json()["id"]
    second = loan(api, add_book(db_session, "Бета"), user=OTHER).json()["id"]
    loan(api, add_book(db_session, "Гамма"))  # на руках
    api.clock.now = day(6)
    act(api, first, "return")
    api.clock.now = day(7)
    api.client.post(f"/library/admin/loans/{second}/return", headers=lib_headers())

    r = api.client.get("/library/admin/export", headers=lib_headers())

    assert r.status_code == 200
    assert "library_history.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    rows = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
    d = lambda n: datetime(2026, 10, n)  # noqa: E731 — openpyxl читает даты как datetime
    assert rows == [
        ("Бета", str(OTHER), d(5), d(12), d(7), "Библиотекарь"),
        ("Альфа", "Иванов Иван", d(5), d(12), d(6), "Пользователь"),
        ("Гамма", "Иванов Иван", d(5), d(12), None, None),
    ]
