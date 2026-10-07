from __future__ import annotations

import threading
import uuid
from datetime import date, datetime, time
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from pa_booking.db.models import AppointmentSlot
from pa_booking.domain.appointments import (
    MSK,
    Kind,
    text_booked,
    text_cancelled_by_user,
    text_slot_removed,
)
from pa_booking.notify.botx import FakeNotifier
from tests.integration.conftest import (
    PSY_ADMIN,
    ApiEnv,
    add_roster,
    bot_headers,
    lk_headers,
)

pytestmark = pytest.mark.db

TODAY = date(2026, 10, 5)  # «сейчас» по умолчанию — 05.10.2026 10:00 МСК
USER = uuid.UUID("11111111-1111-1111-1111-111111111111")
OTHER = uuid.UUID("22222222-2222-2222-2222-222222222222")
PSY = PSY_ADMIN


def at(h: int, m: int = 0, d: date = TODAY) -> datetime:
    return datetime(d.year, d.month, d.day, h, m, tzinfo=MSK)


def add_slot(db: Session, t: time = time(16, 0), d: date = TODAY) -> int:
    slot = AppointmentSlot(slot_date=d, slot_time=t)
    db.add(slot)
    db.commit()
    return slot.id


def add_name(db: Session, huid: uuid.UUID, name: str) -> None:
    add_roster(db, huid, name)


def book(api: ApiEnv, slot_id: int, user: uuid.UUID = USER, kind: str = "psy") -> httpx.Response:
    return api.client.post(
        "/appointments/bookings",
        json={"slot_id": slot_id, "kind": kind},
        headers=bot_headers(user),
    )


def free_ids(api: ApiEnv, d: date = TODAY) -> list[int]:
    r = api.client.get(
        "/appointments/slots",
        params={"from": d.isoformat(), "to": d.isoformat()},
        headers=bot_headers(USER),
    )
    assert r.status_code == 200
    return [s["id"] for day in r.json() for s in day["slots"]]


def my(api: ApiEnv, user: uuid.UUID = USER) -> list[dict[str, object]]:
    r = api.client.get("/appointments/bookings/my", headers=bot_headers(user))
    assert r.status_code == 200
    result: list[dict[str, object]] = r.json()
    return result


def cancel(api: ApiEnv, booking_id: int, user: uuid.UUID = USER) -> httpx.Response:
    return api.client.post(f"/appointments/bookings/{booking_id}/cancel", headers=bot_headers(user))


def admin_delete(api: ApiEnv, slot_id: int) -> httpx.Response:
    return api.client.delete(
        f"/appointments/admin/slots/{slot_id}",
        headers=bot_headers(PSY),
    )


# --- роли ---


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("GET", "/appointments/admin/slots?month=2026-10"),
        ("POST", "/appointments/admin/slots"),
        ("DELETE", "/appointments/admin/slots/1"),
        ("GET", "/appointments/admin/export?all=true"),
    ],
)
@pytest.mark.parametrize("via", ["bot_not_admin", "lk_admin_huid"])
def test_admin_endpoints_require_psychologist_in_bot(
    api: ApiEnv, db_session: Session, method: str, path: str, via: str
) -> None:
    # Админка — только из бота (Р-11): HUID специалиста из ЛК роли не даёт.
    h = (
        bot_headers(USER)
        if via == "bot_not_admin"
        else lk_headers(add_roster(db_session, PSY, "Психолог"))
    )
    r = api.client.request(method, path, json={"date": "2026-10-06", "times": ["16:00"]}, headers=h)
    assert r.status_code == 403
    assert r.json()["code"] == "forbidden"


# --- свободные слоты и запись ---


def test_free_slots_hide_started_booked_and_group_by_date(api: ApiEnv, db_session: Session) -> None:
    add_slot(db_session, time(9, 0))  # уже началось (сейчас 10:00)
    free = add_slot(db_session, time(16, 0))
    booked = add_slot(db_session, time(17, 0))
    tomorrow = add_slot(db_session, time(9, 0), date(2026, 10, 6))
    assert book(api, booked).status_code == 201

    r = api.client.get(
        "/appointments/slots",
        params={"from": "2026-10-05", "to": "2026-10-06"},
        headers=bot_headers(USER),
    )

    assert r.json() == [
        {"date": "2026-10-05", "slots": [{"id": free, "time": "16:00:00"}]},
        {"date": "2026-10-06", "slots": [{"id": tomorrow, "time": "09:00:00"}]},
    ]


def test_booking_notifies_with_roster_name(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    slot_id = add_slot(db_session)

    r = book(api, slot_id, kind="mkr")

    assert r.status_code == 201
    assert r.json() == {
        "id": r.json()["id"],
        "slot_id": slot_id,
        "date": "2026-10-05",
        "time": "16:00:00",
        "kind": "mkr",
        "status": "active",
    }
    assert api.notifier.sent == [text_booked("Иванов Иван", TODAY, time(16, 0), Kind.MKR)]


def test_booking_without_roster_name_uses_huid(api: ApiEnv, db_session: Session) -> None:
    assert book(api, add_slot(db_session)).status_code == 201
    assert api.notifier.sent == [text_booked(str(USER), TODAY, time(16, 0), Kind.PSY)]


def test_fifth_booking_in_month_hits_limit(api: ApiEnv, db_session: Session) -> None:
    for day in (6, 7, 8, 9):
        assert book(api, add_slot(db_session, d=date(2026, 10, day))).status_code == 201

    r = book(api, add_slot(db_session, d=date(2026, 10, 10)))

    assert r.status_code == 409
    assert r.json()["code"] == "monthly_limit"
    assert len(api.notifier.sent) == 4  # отказ — без уведомления
    # Другой месяц — свой лимит.
    assert book(api, add_slot(db_session, d=date(2026, 11, 2))).status_code == 201


def test_booking_taken_or_unknown_slot(api: ApiEnv, db_session: Session) -> None:
    slot_id = add_slot(db_session)
    assert book(api, slot_id, user=OTHER).status_code == 201

    r = book(api, slot_id)
    assert r.status_code == 409
    assert r.json()["code"] == "slot_unavailable"
    assert book(api, 999_999).status_code == 404
    assert len(api.notifier.sent) == 1


def test_parallel_booking_of_one_slot(api: ApiEnv, db_session: Session) -> None:
    slot_id = add_slot(db_session)
    barrier = threading.Barrier(2)
    results: list[httpx.Response] = []

    def worker(user: uuid.UUID) -> None:
        barrier.wait()
        results.append(book(api, slot_id, user=user))

    threads = [threading.Thread(target=worker, args=(u,)) for u in (USER, OTHER)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert sorted(r.status_code for r in results) == [201, 409]
    assert [r.json()["code"] for r in results if r.status_code == 409] == ["slot_unavailable"]


def test_notify_failure_keeps_booking(api: ApiEnv, db_session: Session) -> None:
    api.notifier = FakeNotifier(fail=True)
    assert book(api, add_slot(db_session)).status_code == 201
    assert len(my(api)) == 1


# --- отмена пользователем ---


def test_cancel_foreign_booking_is_404(api: ApiEnv, db_session: Session) -> None:
    booking_id = book(api, add_slot(db_session), user=OTHER).json()["id"]
    r = cancel(api, booking_id)
    assert r.status_code == 404
    assert r.json()["code"] == "not_found"


def test_cancel_before_14_frees_slot(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    slot_id = add_slot(db_session)
    booking_id = book(api, slot_id).json()["id"]
    api.clock.now = at(13, 59)

    r = cancel(api, booking_id)

    assert r.status_code == 200
    assert r.json()["status"] == "cancelled_by_user"
    assert free_ids(api) == [slot_id]
    assert my(api) == []
    assert api.notifier.sent[-1] == text_cancelled_by_user(
        "Иванов Иван", TODAY, time(16, 0), Kind.PSY
    )


def test_cancel_today_after_14_removes_slot(api: ApiEnv, db_session: Session) -> None:
    slot_id = add_slot(db_session)
    booking_id = book(api, slot_id).json()["id"]
    api.clock.now = at(14, 30)

    assert cancel(api, booking_id).status_code == 200
    assert free_ids(api) == []
    # Повторная отмена — активной брони уже нет.
    assert cancel(api, booking_id).status_code == 404


def test_cancel_after_start_is_409(api: ApiEnv, db_session: Session) -> None:
    booking_id = book(api, add_slot(db_session)).json()["id"]
    api.clock.now = at(16, 0)

    r = cancel(api, booking_id)

    assert r.status_code == 409
    assert r.json()["code"] == "too_late_to_cancel"


# --- «мои записи» ---


def test_my_bookings_include_starting_now_and_hide_past(api: ApiEnv, db_session: Session) -> None:
    book(api, add_slot(db_session, time(11, 0)))
    later = book(api, add_slot(db_session, time(16, 0))).json()["id"]
    api.clock.now = at(16, 0)

    assert [b["id"] for b in my(api)] == [later]


# --- специалист ---


def test_admin_add_slots_reports_duplicates(api: ApiEnv, db_session: Session) -> None:
    add_slot(db_session, time(16, 0), date(2026, 10, 6))
    r = api.client.post(
        "/appointments/admin/slots",
        json={"date": "2026-10-06", "times": ["17:00", "16:00", "17:00"]},
        headers=bot_headers(PSY),
    )
    assert r.status_code == 200
    assert r.json() == {"added": ["17:00:00"], "duplicates": ["16:00:00"]}


def test_admin_add_slots_in_past_date_is_422(api: ApiEnv) -> None:
    r = api.client.post(
        "/appointments/admin/slots",
        json={"date": "2026-10-04", "times": ["16:00"]},
        headers=bot_headers(PSY),
    )
    assert r.status_code == 422
    assert r.json()["code"] == "slot_date_in_past"


def test_admin_delete_booked_future_slot_cancels_booking(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    slot_id = add_slot(db_session)
    booking_id = book(api, slot_id).json()["id"]

    r = admin_delete(api, slot_id)

    assert r.status_code == 200
    # Бронь в ответе — боту, чтобы написать человеку лично (Д-6).
    assert r.json() == {
        "cancelled_booking": {
            "id": booking_id,
            "user_huid": str(USER),
            "date": "2026-10-05",
            "time": "16:00:00",
            "kind": "psy",
        }
    }
    assert [b["status"] for b in my(api)] == ["cancelled_by_specialist"]
    assert free_ids(api) == []
    assert api.notifier.sent[-1] == text_slot_removed("Иванов Иван", TODAY, time(16, 0))
    assert admin_delete(api, slot_id).status_code == 404


def test_admin_delete_free_slot_without_notification(api: ApiEnv, db_session: Session) -> None:
    slot_id = add_slot(db_session, time(9, 0))  # уже прошёл — свободный удалить можно
    r = admin_delete(api, slot_id)
    assert (r.status_code, r.json()) == (200, {"cancelled_booking": None})
    assert api.notifier.sent == []


def test_admin_delete_started_booked_slot_is_409(api: ApiEnv, db_session: Session) -> None:
    slot_id = add_slot(db_session)
    book(api, slot_id)
    api.clock.now = at(16, 30)

    r = admin_delete(api, slot_id)

    assert r.status_code == 409
    assert r.json()["code"] == "slot_in_past"


def test_admin_overview_month_with_margin_and_names(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    booked = add_slot(db_session, time(16, 0))
    booking_id = book(api, booked, kind="mkr").json()["id"]
    margin = add_slot(db_session, time(10, 0), date(2026, 11, 3))
    add_slot(db_session, time(10, 0), date(2026, 11, 4))  # вне ± 3 дней

    r = api.client.get(
        "/appointments/admin/slots",
        params={"month": "2026-10"},
        headers=bot_headers(PSY),
    )

    assert r.status_code == 200
    assert r.json() == [
        {
            "date": "2026-10-05",
            "slots": [
                {
                    "id": booked,
                    "time": "16:00:00",
                    "booking": {
                        "id": booking_id,
                        "user_huid": str(USER),
                        "full_name": "Иванов Иван",
                        "kind": "mkr",
                    },
                }
            ],
        },
        {"date": "2026-11-03", "slots": [{"id": margin, "time": "10:00:00", "booking": None}]},
    ]


def test_admin_overview_rejects_bad_month(api: ApiEnv) -> None:
    r = api.client.get(
        "/appointments/admin/slots",
        params={"month": "2026-13"},
        headers=bot_headers(PSY),
    )
    assert r.status_code == 422


# --- выгрузка ---


def export(api: ApiEnv, **params: str) -> httpx.Response:
    return api.client.get(
        "/appointments/admin/export",
        params=params,
        headers=bot_headers(PSY),
    )


def xlsx_rows(r: httpx.Response) -> list[tuple[object, ...]]:
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]


def test_export_month_like_bot_plus_statuses(api: ApiEnv, db_session: Session) -> None:
    add_name(db_session, USER, "Иванов Иван")
    add_slot(db_session, time(11, 0))  # свободный
    assert book(api, add_slot(db_session, time(16, 0)), kind="mkr").status_code == 201
    add_slot(db_session, time(12, 0), date(2026, 11, 2))  # другой месяц
    cancelled = add_slot(db_session, time(17, 0))
    cancel(api, book(api, cancelled, user=OTHER).json()["id"], user=OTHER)

    r = export(api, month="2026-10")

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "appointments_2026-10.xlsx" in r.headers["content-disposition"]
    day = datetime(2026, 10, 5)  # openpyxl читает дату Excel как datetime
    assert xlsx_rows(r) == [
        (day, time(11, 0), None, None, None),
        (day, time(16, 0), "Иванов Иван", "Консультация МКР", "Активна"),
        (day, time(17, 0), str(OTHER), "Психолог", "Отменена пользователем"),
        (day, time(17, 0), None, None, None),  # после отмены до 14:00 слот снова свободен
    ]


def test_export_all_time(api: ApiEnv, db_session: Session) -> None:
    add_slot(db_session, time(11, 0))
    add_slot(db_session, time(12, 0), date(2026, 11, 2))

    r = export(api, all="true")

    assert r.status_code == 200
    assert "appointments_all.xlsx" in r.headers["content-disposition"]
    assert len(xlsx_rows(r)) == 2


@pytest.mark.parametrize("params", [{}, {"all": "false"}, {"all": "true", "month": "2026-10"}])
def test_export_needs_exactly_one_period(api: ApiEnv, params: dict[str, str]) -> None:
    r = export(api, **params)
    assert r.status_code == 422
    assert r.json()["code"] == "validation_error"
